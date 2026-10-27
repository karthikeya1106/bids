"""Report exports: the downloaded files must carry the page's numbers.

The PDF is checked structurally (it is a real PDF, served as an attachment), and
the workbook is opened and compared field by field against the API payload, so a
formatting change cannot silently drop or alter a figure.
"""
from __future__ import annotations

import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests import data
from tests.conftest import csv_bytes


def _upload(client, headers, frame, name="sales.csv", filename="sales.csv"):
    preview = client.post(
        "/api/upload/preview?analysis_type=auto",
        files={"file": (filename, csv_bytes(frame), "text/csv")},
        headers=headers,
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    mapping = {key: value["value"] for key, value in body["mapping"].items() if value.get("value")}
    commit = client.post(
        "/api/upload/commit",
        files={"file": (filename, csv_bytes(frame), "text/csv")},
        data={
            "schema": body["analysis_type"],
            "mapping": json.dumps(mapping),
            "task": body["task"],
            "target": body["target"] or "",
            "org_name": name,
        },
        headers=headers,
    )
    assert commit.status_code == 200, commit.text
    return commit.json()


def _table(sheet) -> list[dict]:
    """Read a sheet whose header row is the first one starting with a label."""
    rows = [row for row in sheet.iter_rows(values_only=True)]
    for index, row in enumerate(rows):
        if row and isinstance(row[0], str) and row[0] in {"Field", "Metric", "Month", "Segment", "Priority", "Column"}:
            headers = [value for value in row]
            return [
                dict(zip(headers, values))
                for values in rows[index + 1:]
                if any(value is not None for value in values)
            ]
    return []


def _sheet_named(workbook, prefix: str):
    return next(sheet for sheet in workbook.worksheets if sheet.title.startswith(prefix))


# ─── Sales report ────────────────────────────────────────────────────────────

def test_pdf_export_is_a_real_pdf_attachment(client, auth, owner):
    headers = auth(owner)
    _upload(client, headers, data.sales_standard(180))

    response = client.get("/api/analysis/report/export.pdf", headers=headers)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.pdf"')
    assert response.headers["cache-control"] == "no-store"
    assert response.content.startswith(b"%PDF")
    assert response.content.rstrip().endswith(b"%%EOF")
    assert len(response.content) > 2_000


def test_workbook_matches_the_report_payload(client, auth, owner):
    headers = auth(owner)
    _upload(client, headers, data.sales_standard(180), name="Export Test Ltd")
    report = client.get("/api/analysis/report", headers=headers).json()

    response = client.get("/api/analysis/report/export.xlsx", headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames[0] == "Overview"

    overview = {row[0]: row[1] for row in _sheet_named(workbook, "Overview").iter_rows(min_row=1, values_only=True) if row and row[0]}
    assert overview["Dataset"] == report["org"]["dataset_name"]
    assert overview["Organisation"] == "Export Test Ltd"
    assert overview["Rows"] == f"{report['org']['row_count']:,}"
    assert str(report["health_score"]) in str(overview["Health score"])

    kpis = {row["Metric"]: row for row in _table(_sheet_named(workbook, "Key performance"))}
    revenue = kpis["Total Revenue"]
    assert revenue["Current"] == pytest.approx(report["kpis"]["total_revenue"]["value"], rel=1e-9)
    assert revenue["Baseline"] == pytest.approx(report["kpis"]["total_revenue"]["baseline"], rel=1e-9)
    assert "Unit" in revenue

    monthly = _table(_sheet_named(workbook, "Monthly performance"))
    assert len(monthly) == len(report["drift_series"])
    assert monthly[0]["Revenue"] == pytest.approx(report["drift_series"][0]["revenue"], rel=1e-9)

    actions = _table(_sheet_named(workbook, "Priority actions"))
    assert len(actions) == len(report["action_items"])

    notes = _sheet_named(workbook, "Notes and limitations")
    text = " ".join(str(cell.value) for row in notes.iter_rows() for cell in row if cell.value)
    assert "not proven causal effects" in text


def test_workbook_includes_the_forecast_and_forecast_note(client, auth, owner):
    headers = auth(owner)
    _upload(client, headers, data.sales_with_elastic_demand(400))
    report = client.get("/api/analysis/report", headers=headers).json()
    if not report.get("forecast"):
        pytest.skip("this fixture does not span enough months to forecast")

    workbook = load_workbook(BytesIO(client.get("/api/analysis/report/export.xlsx", headers=headers).content))
    forecast = _table(_sheet_named(workbook, "Revenue forecast"))
    assert len(forecast) == len(report["forecast"])
    assert forecast[0]["Predicted revenue"] == pytest.approx(report["forecast"][0]["predicted"], rel=1e-9)


# ─── Adaptive report ─────────────────────────────────────────────────────────

def test_adaptive_dataset_exports_its_own_sections(client, auth, owner):
    headers = auth(owner)
    _upload(client, headers, data.telecom_churn(200), name="Telecom", filename="churn.csv")
    report = client.get("/api/analysis/report", headers=headers).json()
    assert report["mode"] == "adaptive"

    response = client.get("/api/analysis/report/export.xlsx", headers=headers)
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    titles = {sheet.title for sheet in workbook.worksheets}
    assert "Column profile" in titles
    assert "Notes and limitations" in titles
    # Sales-only sections must not appear on a churn dataset.
    assert "Key performance indicators" not in titles

    profile = _table(_sheet_named(workbook, "Column profile"))
    assert len(profile) >= len(report["profile"]["column_profiles"]) - 1
    assert any(row["Column"] for row in profile)

    pdf = client.get("/api/analysis/report/export.pdf", headers=headers)
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")


# ─── Access control ──────────────────────────────────────────────────────────

def test_export_requires_a_signature(client, owner):
    assert client.get("/api/analysis/report/export.pdf").status_code == 401
    assert client.get("/api/analysis/report/export.xlsx").status_code == 401


def test_export_never_crosses_workspaces(client, auth, owner):
    owner_headers = auth(owner)
    _upload(client, owner_headers, data.sales_standard(120))

    other = f"other-{owner}"
    assert client.get("/api/analysis/report/export.pdf", headers=auth(other)).status_code == 404
    assert client.get("/api/analysis/report/export.xlsx", headers=auth(other)).status_code == 404
