"""
PDF and Excel exports of the analysis report.

Both formats are rendered from one section model built from the payload that
``GET /api/analysis/report`` returns, so a downloaded file never disagrees with
the page. Nothing here recomputes analytics: it only formats what the modules
already produced, and it keeps the report's limitations in the file.
"""
from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from services.semantics import humanize

MAX_PDF_ROWS = 200

MONEY = "money"
PERCENT = "percent"
NUMBER = "number"
DECIMAL = "decimal"
TEXT = "text"

SALES_LIMITATIONS = [
    "Every figure is aggregated from the rows this workspace uploaded; no external data is blended in.",
    "The health score weights revenue trend (35), margin trend (25), customer health (20) and an anomaly penalty (20).",
    "Root-cause factors are statistical associations in this dataset, not proven causal effects.",
    "The forecast is a linear trend with a residual band; it does not model seasonality, promotion or competitor response.",
    "Benchmark ranges are quartiles of this dataset or stated retail priors, not an external industry survey.",
]

# ─── Formatting ───────────────────────────────────────────────────────────────


def _number(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if result != result:  # NaN
        return None
    return result


def _safe(text: Any) -> str:
    """Restrict text to what the PDF's built-in fonts can draw."""
    cleaned = (
        str(text)
        .replace("\u2014", "-")
        .replace("\u2013", "-")
        .replace("\u2019", "'")
        .replace("\u2018", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2192", "->")
        .replace("\u2022", "-")
        .replace("\u00a0", " ")
    )
    return cleaned.encode("latin-1", "replace").decode("latin-1")


def _fmt(value: Any, kind: str) -> str:
    if kind == TEXT:
        return "" if value is None else str(value)
    number = _number(value)
    if number is None:
        return "-"
    if kind == MONEY:
        return f"${number:,.2f}" if abs(number) < 1000 else f"${number:,.0f}"
    if kind == PERCENT:
        return f"{number:,.2f}%"
    if kind == NUMBER:
        return f"{number:,.0f}"
    if abs(number) < 1:
        return f"{number:,.3f}"
    return f"{number:,.2f}"


def _change(value: Any, baseline: Any) -> Optional[float]:
    current, base = _number(value), _number(baseline)
    if current is None or base in (None, 0):
        return None
    return (current - base) / abs(base) * 100


def _label(key: str) -> str:
    return humanize(key).replace("_", " ").strip().title() if "_" in key else humanize(key)


# ─── Section model ────────────────────────────────────────────────────────────


def _column(key: str, label: str, kind: str = TEXT) -> dict:
    return {"key": key, "label": label, "kind": kind}


def _section(key: str, title: str, columns: list[dict], rows: Iterable[dict], note: str = "") -> dict:
    return {"key": key, "title": title, "columns": columns, "rows": list(rows), "note": note}


def _overview(pairs: list[tuple[str, str]]) -> dict:
    return _section(
        "overview",
        "Overview",
        [_column("label", "Field"), _column("value", "Value")],
        [{"label": label, "value": str(value)} for label, value in pairs],
    )


def _sales_sections(report: dict) -> list[dict]:
    org = report.get("org") or {}
    sections: list[dict] = [
        _overview([
            ("Dataset", org.get("dataset_name") or "-"),
            ("Organisation", org.get("org_name") or "-"),
            ("Dataset type", report.get("schema_label") or "Sales & Revenue Analytics"),
            ("Rows", f"{int(org.get('row_count') or 0):,}"),
            ("Customers", f"{int(org.get('customer_count') or 0):,}"),
            ("Total sales", _fmt(org.get("total_sales"), MONEY)),
            ("Period", f"{org.get('date_from') or '?'} to {org.get('date_to') or '?'}"),
            ("Health score", _fmt(report.get("health_score"), DECIMAL) + " / 100"),
            ("Generated", str(report.get("generated_at") or "").replace("T", " ")[:19]),
        ])
    ]

    kpis = report.get("kpis") or {}
    if kpis:
        sections.append(_section(
            "kpis",
            "Key performance indicators",
            [
                _column("label", "Metric"),
                _column("value", "Current", DECIMAL),
                _column("baseline", "Baseline", DECIMAL),
                _column("change", "Change vs baseline (%)", DECIMAL),
                _column("target", "Target", DECIMAL),
                _column("unit", "Unit"),
            ],
            [
                {
                    "label": _label(key),
                    "value": entry.get("value"),
                    "baseline": entry.get("baseline"),
                    "change": _change(entry.get("value"), entry.get("baseline")),
                    "target": entry.get("target"),
                    "unit": entry.get("unit"),
                }
                for key, entry in kpis.items()
            ],
            note="The baseline is the first half of the dataset's period; the current value covers the second half.",
        ))

    segmentation = report.get("segmentation") or []
    if segmentation:
        sections.append(_section(
            "segments",
            "Customer segments (RFM)",
            [
                _column("segment", "Segment"),
                _column("customer_count", "Customers", NUMBER),
                _column("avg_spend", "Avg spend", MONEY),
                _column("avg_frequency", "Avg orders", DECIMAL),
                _column("avg_recency_days", "Avg days since last order", NUMBER),
            ],
            segmentation,
            note="Segments are assigned from recency, frequency and monetary percentiles within this dataset.",
        ))

    monthly = report.get("drift_series") or []
    if monthly:
        sections.append(_section(
            "monthly",
            "Monthly performance",
            [
                _column("month", "Month"),
                _column("revenue", "Revenue", MONEY),
                _column("profit", "Profit", MONEY),
                _column("orders", "Orders", NUMBER),
                _column("customers", "Customers", NUMBER),
                _column("aov", "Avg order value", MONEY),
                _column("margin_pct", "Margin (%)", DECIMAL),
            ],
            monthly,
        ))

    root_cause = report.get("root_cause") or {}
    fishbone = root_cause.get("fishbone") or {}
    cause_rows = [
        {
            "branch": branch,
            "factor": cause.get("factor"),
            "confidence": (_number(cause.get("confidence")) or 0) * 100,
            "evidence": cause.get("evidence"),
        }
        for branch, causes in fishbone.items()
        for cause in causes or []
    ]
    if cause_rows:
        sections.append(_section(
            "root-cause",
            "Root cause analysis",
            [
                _column("branch", "Branch"),
                _column("factor", "Factor"),
                _column("confidence", "Confidence (%)", DECIMAL),
                _column("evidence", "Evidence"),
            ],
            cause_rows,
            note="Confidence is the strength of the association in this data, not a causal probability.",
        ))

    anomalies = report.get("anomalies") or []
    if anomalies:
        sections.append(_section(
            "anomalies",
            "Detected anomalies",
            [
                _column("date", "Period"),
                _column("type", "Type"),
                _column("value", "Observed", DECIMAL),
                _column("z_score", "Z-score", DECIMAL),
            ],
            [
                {
                    "date": item.get("date") or item.get("month"),
                    "type": item.get("type"),
                    "value": item.get("value"),
                    "z_score": item.get("z_score"),
                }
                for item in anomalies
            ],
            note="Anomalies are periods more than two standard deviations from the fitted trend of the monthly revenue series.",
        ))

    forecast = report.get("forecast") or []
    if forecast:
        sections.append(_section(
            "forecast",
            "Revenue forecast",
            [
                _column("month", "Month"),
                _column("predicted", "Predicted revenue", MONEY),
                _column("lower", "Lower bound", MONEY),
                _column("upper", "Upper bound", MONEY),
            ],
            forecast,
            note="Bounds are a 95% band from the residual spread of the fitted linear trend.",
        ))

    benchmarks = report.get("benchmarks") or []
    if benchmarks:
        sections.append(_section(
            "benchmarks",
            "Benchmark comparison",
            [
                _column("metric", "Metric"),
                _column("org_value", "This dataset", DECIMAL),
                _column("benchmark_low", "Range low", DECIMAL),
                _column("benchmark_high", "Range high", DECIMAL),
                _column("unit", "Unit"),
            ],
            benchmarks,
        ))

    actions = report.get("action_items") or []
    if actions:
        sections.append(_section(
            "actions",
            "Priority actions",
            [
                _column("priority", "Priority", NUMBER),
                _column("area", "Area"),
                _column("action", "Action"),
                _column("impact", "Impact"),
                _column("effort", "Effort"),
            ],
            actions,
        ))

    return sections


def _adaptive_sections(report: dict) -> list[dict]:
    org = report.get("org") or {}
    profile = report.get("profile") or {}
    adaptive = report.get("adaptive") or {}
    capabilities = report.get("capabilities") or {}
    summary = adaptive.get("summary") or {}
    time_coverage = profile.get("time_coverage") or {}

    sections: list[dict] = [
        _overview([
            ("Dataset", org.get("dataset_name") or "-"),
            ("Organisation", org.get("org_name") or "-"),
            ("Dataset type", report.get("schema_label") or report.get("dataset_type") or "-"),
            ("Analysis task", str(report.get("task") or summary.get("task") or "eda").replace("_", " ").title()),
            ("Target column", report.get("target") or "none"),
            ("Rows", f"{int(org.get('row_count') or summary.get('rows') or 0):,}"),
            ("Columns", f"{int(summary.get('columns') or profile.get('columns') or 0):,}"),
            ("Missing values", f"{_fmt(profile.get('missing_pct'), DECIMAL)} of cells"),
            ("Duplicate rows", f"{int(profile.get('duplicate_rows') or 0):,}"),
            ("Period", f"{org.get('date_from') or time_coverage.get('from') or 'no date column'} to {org.get('date_to') or time_coverage.get('to') or ''}".strip()),
            ("Generated", str(report.get("generated_at") or adaptive.get("generated_at") or "").replace("T", " ")[:19]),
        ])
    ]

    columns = profile.get("column_profiles") or []
    if columns:
        sections.append(_section(
            "columns",
            "Column profile",
            [
                _column("name", "Column"),
                _column("dtype", "Stored type"),
                _column("missing_pct", "Missing (%)", DECIMAL),
                _column("unique", "Unique values", NUMBER),
                _column("mean", "Mean", DECIMAL),
                _column("min", "Min", DECIMAL),
                _column("max", "Max", DECIMAL),
            ],
            columns,
            note="Means, minima and maxima are shown only for columns where they are meaningful.",
        ))

    quality = profile.get("quality") or []
    if quality:
        sections.append(_section(
            "quality",
            "Data quality findings",
            [
                _column("level", "Level"),
                _column("column", "Column"),
                _column("message", "Finding"),
            ],
            quality,
            note="Findings are observations about this file; they do not change the stored data.",
        ))

    correlations = profile.get("correlations") or []
    if correlations:
        sections.append(_section(
            "correlations",
            "Strongest correlations",
            [
                _column("a", "Column A"),
                _column("b", "Column B"),
                _column("r", "Pearson r", DECIMAL),
                _column("strength", "Strength"),
            ],
            correlations,
            note="Correlation measures association. It does not establish that one column causes the other.",
        ))

    modules = capabilities.get("modules") or {}
    if modules:
        sections.append(_section(
            "analyses",
            "Supported analyses",
            [
                _column("module", "Module"),
                _column("status", "Status"),
                _column("reason", "Why"),
            ],
            [
                {
                    "module": _label(name),
                    "status": (entry or {}).get("status") or "unknown",
                    "reason": "; ".join((entry or {}).get("reasons") or []) or "Available on this dataset.",
                }
                for name, entry in modules.items()
                if name != "analysis"
            ],
            note="A module runs only when this dataset contains what it needs; unavailable modules say what is missing.",
        ))

    task_result = adaptive.get("task") or {}
    metrics = task_result.get("metrics") or {}
    if metrics:
        sections.append(_section(
            "model",
            f"Modelled task: {str(task_result.get('type') or report.get('task') or 'model').replace('_', ' ').title()}",
            [_column("label", "Metric"), _column("value", "Value")],
            [
                {"label": _label(key), "value": _fmt(value, DECIMAL) if _number(value) is not None else str(value)}
                for key, value in metrics.items()
            ],
            note="Metrics are measured on a held-out split of this file or on in-sample fit where stated. "
                 + (str(task_result.get("method") or "") if task_result.get("method") else ""),
        ))

    clusters = adaptive.get("clusters") or {}
    cluster_profiles = clusters.get("profiles") or []
    if cluster_profiles:
        keys = list(cluster_profiles[0].keys())[:8]
        sections.append(_section(
            "clusters",
            "Cluster profiles",
            [_column(key, _label(key), DECIMAL if _number(cluster_profiles[0].get(key)) is not None else TEXT) for key in keys],
            cluster_profiles,
            note="Cluster labels are ordered by size and are descriptive, not business names.",
        ))

    return sections


def _limitations(report: dict) -> list[str]:
    if report.get("mode") == "sales":
        limitations = list(SALES_LIMITATIONS)
        notes = report.get("notes") or []
    else:
        adaptive = report.get("adaptive") or {}
        limitations = list(adaptive.get("limitations") or [])
        notes = list(report.get("notes") or [])
    limitations.extend(note for note in notes if note)
    return limitations[:20]


def report_sections(report: dict) -> tuple[dict, list[dict], list[str]]:
    """Build (header, sections, limitations) for one report payload."""
    org = report.get("org") or {}
    header = {
        "title": "BIDS Intelligence - Analysis Report",
        "dataset": org.get("dataset_name") or "Dataset",
        "org": org.get("org_name") or "",
        "generated_at": str(report.get("generated_at") or "").replace("T", " ")[:19],
        "schema_label": report.get("schema_label") or "",
        "mode": report.get("mode") or "sales",
    }
    sections = _sales_sections(report) if header["mode"] == "sales" else _adaptive_sections(report)
    return header, sections, _limitations(report)


def report_filename(report: dict, extension: str) -> str:
    name = str((report.get("org") or {}).get("dataset_name") or "dataset")
    slug = re.sub(r"[^a-z0-9]+", "-", name.rsplit(".", 1)[0].lower()).strip("-") or "dataset"
    stamp = str(report.get("generated_at") or "")[:10].replace("-", "") or "report"
    return f"bids-report-{slug}-{stamp}.{extension}"


# ─── PDF ──────────────────────────────────────────────────────────────────────


def _pdf_styles() -> dict:
    return {
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.black),
        "meta": ParagraphStyle("meta", fontName="Helvetica", fontSize=8.5, leading=12, textColor=colors.HexColor("#4b5563")),
        "section": ParagraphStyle("section", fontName="Helvetica-Bold", fontSize=10.5, leading=14, textColor=colors.black, spaceBefore=4),
        "note": ParagraphStyle("note", fontName="Helvetica-Oblique", fontSize=7, leading=9, textColor=colors.HexColor("#6b7280")),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=7, leading=9, textColor=colors.HexColor("#111827")),
        "head": ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=7, leading=9, textColor=colors.white),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=8, leading=11.5, textColor=colors.HexColor("#111827")),
    }


def _pdf_cell(value: Any, kind: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(_safe(_fmt(value, kind) if kind != TEXT else value), style)


def _column_widths(columns: list[dict], rows: list[dict], total: float) -> list[float]:
    weights = []
    for column in columns:
        longest = max([len(str(column["label"]))] + [len(_fmt(row.get(column["key"]), column["kind"])) for row in rows[:40]])
        weights.append(max(8, min(52, longest)))
    scale = total / sum(weights)
    return [weight * scale for weight in weights]


def _pdf_section(section: dict, styles: dict, width: float) -> list:
    story: list = [Paragraph(_safe(section["title"]), styles["section"]), Spacer(1, 3)]
    if section.get("note"):
        story.append(Paragraph(_safe(section["note"]), styles["note"]))
        story.append(Spacer(1, 3))

    rows = section["rows"][:MAX_PDF_ROWS]
    header = [Paragraph(_safe(column["label"]), styles["head"]) for column in section["columns"]]
    data = [header] + [
        [_pdf_cell(row.get(column["key"]), column["kind"], styles["cell"]) for column in section["columns"]]
        for row in rows
    ]
    widths = _column_widths(section["columns"], rows, width)
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.black),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]
    for index in range(2, len(data), 2):
        style.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#f9fafb")))
    table.setStyle(TableStyle(style))
    story.append(table)
    if len(section["rows"]) > MAX_PDF_ROWS:
        story.append(Spacer(1, 2))
        story.append(Paragraph(
            _safe(f"Showing the first {MAX_PDF_ROWS:,} of {len(section['rows']):,} rows; the Excel export has all of them."),
            styles["note"],
        ))
    story.append(Spacer(1, 9))
    return story


def build_pdf(report: dict) -> bytes:
    """Render the report as a landscape A4 PDF."""
    header, sections, limitations = report_sections(report)
    styles = _pdf_styles()
    buffer = BytesIO()
    page_width, page_height = landscape(A4)
    margin = 14 * mm
    width = page_width - 2 * margin

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#e5e7eb"))
        canvas.setLineWidth(0.5)
        canvas.line(margin, page_height - margin + 4 * mm, page_width - margin, page_height - margin + 4 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawString(margin, 8 * mm, _safe(f"{header['title']} | {header['dataset']}"))
        canvas.drawRightString(page_width - margin, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=13 * mm,
        title=f"{header['title']} - {header['dataset']}",
        author="BIDS Intelligence",
        subject=header["schema_label"] or "Analysis report",
    )

    story: list = [Paragraph(_safe(header["title"]), styles["title"]), Spacer(1, 3)]
    meta = " | ".join(part for part in (header["dataset"], header["org"], header["schema_label"], f"Generated {header['generated_at']}") if part)
    story.append(Paragraph(_safe(meta), styles["meta"]))
    story.append(Spacer(1, 7))
    story.append(Paragraph(
        _safe("This document reports results computed from the dataset this workspace uploaded. "
              "Unsupported analyses are named as limitations rather than estimated."),
        styles["body"],
    ))
    story.append(Spacer(1, 10))

    for section in sections:
        story.extend(_pdf_section(section, styles, width))

    story.append(Paragraph("Limitations and notes", styles["section"]))
    story.append(Spacer(1, 3))
    for line in limitations:
        story.append(Paragraph(_safe(f"- {line}"), styles["body"]))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        _safe("Numbers are only as reliable as the uploaded file. Re-running an analysis after a new upload can change them."),
        styles["note"],
    ))

    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


# ─── Excel ────────────────────────────────────────────────────────────────────

_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
_NUMBER_FORMATS = {
    MONEY: '"$"#,##0.00',
    PERCENT: '0.00"%"',
    NUMBER: "#,##0",
    DECIMAL: "#,##0.00####",
}


def _sheet_title(title: str, used: set[str]) -> str:
    base = _SHEET_FORBIDDEN.sub(" ", title).strip() or "Sheet"
    base = base[:31]
    candidate, index = base, 2
    while candidate.lower() in used:
        suffix = f" {index}"
        candidate = base[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.lower())
    return candidate


def _write_sheet(sheet, section: dict) -> None:
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet["A1"] = section["title"]
    sheet["A1"].font = Font(bold=True, size=13)
    row = 2
    if section.get("note"):
        sheet.cell(row=2, column=1, value=section["note"]).font = Font(italic=True, size=9, color="6B7280")
        row = 3
    header_row = row + 1

    for index, column in enumerate(section["columns"], start=1):
        cell = sheet.cell(row=header_row, column=index, value=column["label"])
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="111827")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border

    for offset, entry in enumerate(section["rows"], start=1):
        for index, column in enumerate(section["columns"], start=1):
            raw = entry.get(column["key"])
            kind = column["kind"]
            cell = sheet.cell(row=header_row + offset, column=index)
            if kind == TEXT:
                cell.value = "" if raw is None else str(raw)
            else:
                cell.value = _number(raw)
                if cell.value is not None:
                    cell.number_format = _NUMBER_FORMATS.get(kind, "General")
            cell.alignment = Alignment(vertical="top", wrap_text=kind == TEXT)
            cell.border = border

    for index, column in enumerate(section["columns"], start=1):
        longest = max(
            [len(column["label"])]
            + [len(_fmt(entry.get(column["key"]), column["kind"])) for entry in section["rows"][:200]]
        )
        bracket = sheet.cell(row=header_row, column=index).coordinate
        sheet.column_dimensions[get_column_letter(index)].width = max(10, min(46, longest + 2))
        del bracket

    last_row = header_row + len(section["rows"])
    if section["rows"]:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(section['columns']))}{last_row}"
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.sheet_view.showGridLines = False


def build_xlsx(report: dict) -> bytes:
    """Render the report as an Excel workbook: one sheet per section plus notes."""
    header, sections, limitations = report_sections(report)
    workbook = Workbook()
    overview = workbook.active
    used: set[str] = set()

    first = sections[0] if sections else _overview([("Dataset", header["dataset"])])
    overview.title = _sheet_title(first["title"], used)
    _write_sheet(overview, first)

    for section in sections[1:]:
        sheet = workbook.create_sheet(_sheet_title(section["title"], used))
        _write_sheet(sheet, section)

    notes = workbook.create_sheet(_sheet_title("Notes and limitations", used))
    notes["A1"] = "Notes and limitations"
    notes["A1"].font = Font(bold=True, size=13)
    notes["A2"] = f"{header['dataset']} - generated {header['generated_at']}"
    notes["A2"].font = Font(italic=True, size=9, color="6B7280")
    for offset, line in enumerate(limitations, start=4):
        notes.cell(row=offset, column=1, value=f"- {line}")
        notes.cell(row=offset, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    notes.column_dimensions["A"].width = 120
    notes.sheet_view.showGridLines = False

    workbook.properties.title = f"{header['title']} - {header['dataset']}"
    workbook.properties.creator = "BIDS Intelligence"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
